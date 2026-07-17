#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
import re
import zipfile
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
import pandas as pd


DEFAULT_SUPPLEMENT = Path("data/raw/proteus_2025/mic-171-01580-s001.xlsx")
MANUAL_FALLBACKS = [
    Path("~/Downloads/mic-171-01580-s001.xlsx").expanduser(),
    Path("data/raw/proteus_2025/Table_S1.xlsx"),
    Path("data/raw/proteus_2025/proteus_table_s1.xlsx"),
]


def normalize_name(value: object) -> str:
    return re.sub(r"[^a-z0-9]+", "_", str(value).strip().lower()).strip("_")


def is_valid_xlsx(path: Path) -> bool:
    return path.exists() and path.stat().st_size > 1_000 and zipfile.is_zipfile(path)


def resolve_supplement(user_path: str | None) -> Path | None:
    if user_path:
        path = Path(user_path).expanduser()
        return path if is_valid_xlsx(path) else None
    candidates = []
    candidates.append(DEFAULT_SUPPLEMENT)
    candidates.extend(MANUAL_FALLBACKS)
    for path in candidates:
        if is_valid_xlsx(path):
            return path
    return None


def infer_header_row(raw: pd.DataFrame) -> int:
    tokens = (
        "isolate",
        "strain",
        "sample",
        "accession",
        "biosample",
        "sra",
        "bioproject",
        "chlorhexidine",
        "chd",
        "mic",
        "octenidine",
        "cationic",
        "biocide",
    )
    best_index = 0
    best_score = -1
    for index in range(min(30, len(raw))):
        values = [normalize_name(value) for value in raw.iloc[index].tolist()]
        score = sum(any(token in value for value in values) for token in tokens)
        non_empty = sum(value not in {"", "nan", "none"} for value in values)
        score = score * 10 + min(non_empty, 20)
        if score > best_score:
            best_score = score
            best_index = index
    return best_index


def read_raw_sheet_preview(path: Path, sheet_name: str, max_rows: int | None = None) -> pd.DataFrame:
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook[sheet_name]
    rows: list[list[object]] = []
    for index, row in enumerate(worksheet.iter_rows(values_only=True)):
        if max_rows is not None and index >= max_rows:
            break
        rows.append(list(row))
    workbook.close()
    return pd.DataFrame(rows)


def read_normalized_sheet(path: Path, sheet_name: str, max_rows: int | None = None) -> pd.DataFrame:
    raw = read_raw_sheet_preview(path, sheet_name, max_rows=max_rows)
    if raw.empty:
        return pd.DataFrame()
    header_index = infer_header_row(raw)
    columns = [str(value).strip() if str(value).strip() else f"unnamed_{i}" for i, value in enumerate(raw.iloc[header_index])]
    frame = raw.iloc[header_index + 1 :].copy()
    frame.columns = columns
    frame = frame.dropna(how="all").reset_index(drop=True)
    frame = frame.loc[:, ~pd.Index([normalize_name(column) for column in frame.columns]).duplicated()]
    return frame


def score_frame(frame: pd.DataFrame) -> int:
    names = [normalize_name(column) for column in frame.columns]
    tokens = ("isolate", "strain", "accession", "biosample", "sra", "bioproject", "chlorhexidine", "chd", "mic")
    return sum(any(token in name for name in names) for token in tokens) + min(len(frame), 200) // 20


def inventory_workbook(path: Path) -> tuple[pd.DataFrame, dict[str, pd.DataFrame]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheets: dict[str, pd.DataFrame] = {}
    rows: list[dict[str, Any]] = []
    preferred = [
        sheet_name
        for sheet_name in workbook.sheetnames
        if normalize_name(sheet_name) in {"isolate_origins", "mutations_in_adaptations", "snippy_files"}
    ]
    sheet_names = preferred or workbook.sheetnames[:5]
    for sheet_name in sheet_names:
        worksheet = workbook[sheet_name]
        frame = read_normalized_sheet(path, sheet_name, max_rows=500)
        sheets[sheet_name] = frame
        columns = list(frame.columns)
        normalized = [normalize_name(column) for column in columns]
        rows.append(
            {
                "sheet_name": sheet_name,
                "rows": int(len(frame)),
                "columns": int(len(columns)),
                "score": int(score_frame(frame)),
                "mentions_isolate": any("isolate" in name or "strain" in name for name in normalized),
                "mentions_accession": any("accession" in name or "biosample" in name or "sra" in name or "bioproject" in name for name in normalized),
                "mentions_chx_or_mic": any("chlorhexidine" in name or "chd" in name or "mic" in name for name in normalized),
                "column_preview": "; ".join(columns[:25]),
            }
        )
    workbook.close()
    return pd.DataFrame(rows).sort_values(["score", "rows"], ascending=[False, False]), sheets


def find_column(frame: pd.DataFrame, candidates: list[str], required: bool = False) -> str | None:
    normalized = {normalize_name(column): column for column in frame.columns}
    for candidate in candidates:
        key = normalize_name(candidate)
        if key in normalized:
            return normalized[key]
    for key, column in normalized.items():
        if any(normalize_name(candidate) in key for candidate in candidates):
            return column
    if required:
        raise KeyError(f"Could not find any of {candidates}. Available columns: {list(frame.columns)}")
    return None


def candidate_accession_columns(frame: pd.DataFrame) -> list[str]:
    keep = []
    for column in frame.columns:
        name = normalize_name(column)
        if any(token in name for token in ["accession", "biosample", "sra", "bioproject", "run"]):
            keep.append(column)
    return keep


def candidate_phenotype_columns(frame: pd.DataFrame) -> list[str]:
    keep = []
    for column in frame.columns:
        name = normalize_name(column)
        if any(token in name for token in ["mic", "chlorhexidine", "chd", "octenidine", "oct", "cationic", "biocide"]):
            if not any(skip in name for skip in ["reference", "method", "comment", "note"]):
                keep.append(column)
    return keep


def extract_numeric_like(value: object) -> str:
    if pd.isna(value):
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def build_manifest(frame: pd.DataFrame, source_sheet: str) -> tuple[pd.DataFrame, pd.DataFrame]:
    isolate_col = find_column(frame, ["isolate", "isolate_id", "strain", "strain_id", "sample", "sample_id"], required=True)
    accession_cols = candidate_accession_columns(frame)
    phenotype_cols = candidate_phenotype_columns(frame)

    base_records: list[dict[str, Any]] = []
    phenotype_records: list[dict[str, Any]] = []
    for _, row in frame.iterrows():
        isolate_id = extract_numeric_like(row.get(isolate_col))
        if not isolate_id or isolate_id.lower() in {"nan", "none"}:
            continue
        accessions = {normalize_name(column): extract_numeric_like(row.get(column)) for column in accession_cols}
        base = {
            "study_key": "proteus_2025_chx_cationic",
            "source_sheet": source_sheet,
            "organism": "Proteus mirabilis",
            "isolate_id": isolate_id,
            "doi": "10.1099/mic.0.001580",
            "source_url": "https://pmc.ncbi.nlm.nih.gov/articles/PMC12304739/",
            **accessions,
        }
        base_records.append(base)
        for column in phenotype_cols:
            value = extract_numeric_like(row.get(column))
            if not value or value.lower() in {"nan", "none", "not determined", "nd"}:
                continue
            phenotype_records.append(
                {
                    **base,
                    "compound_or_endpoint": str(column),
                    "phenotype_value": value,
                    "phenotype_unit": "",
                    "phenotype_type": "reported_table_value",
                }
            )

    base_manifest = pd.DataFrame(base_records).drop_duplicates()
    phenotype_manifest = pd.DataFrame(phenotype_records).drop_duplicates()
    return base_manifest, phenotype_manifest


def empty_outputs(args: argparse.Namespace, report: dict[str, Any]) -> None:
    Path(args.sheet_inventory).parent.mkdir(parents=True, exist_ok=True)
    Path(args.report_output).parent.mkdir(parents=True, exist_ok=True)
    pd.DataFrame().to_csv(args.sheet_inventory, index=False)
    pd.DataFrame().to_csv(args.isolate_output, index=False)
    pd.DataFrame().to_csv(args.phenotype_output, index=False)
    Path(args.report_output).write_text(json.dumps(report, indent=2) + "\n")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--supplement", default=None)
    parser.add_argument("--sheet-inventory", default="data/manifests/proteus_2025_table_s1_sheet_inventory.csv")
    parser.add_argument("--isolate-output", default="data/manifests/proteus_2025_candidate_isolate_manifest.csv")
    parser.add_argument("--phenotype-output", default="data/manifests/proteus_2025_candidate_phenotype_manifest.csv")
    parser.add_argument("--report-output", default="reports/proteus_2025_manifest_extraction.json")
    args = parser.parse_args()

    supplement = resolve_supplement(args.supplement)
    if supplement is None:
        report = {
            "decision": "blocked_manual_supplement_download_required",
            "valid_supplement_found": False,
            "expected_file": str(DEFAULT_SUPPLEMENT),
            "manual_fallbacks_checked": [str(path) for path in MANUAL_FALLBACKS],
            "recommended_next_action": (
                "Download mic-171-01580-s001.xlsx from the PMC article in a browser, place it at "
                "data/raw/proteus_2025/mic-171-01580-s001.xlsx or ~/Downloads/mic-171-01580-s001.xlsx, then rerun this script."
            ),
        }
        empty_outputs(args, report)
        print(report["decision"])
        return

    inventory, sheets = inventory_workbook(supplement)
    best_sheet = str(inventory.iloc[0]["sheet_name"])
    best_frame = sheets[best_sheet]
    try:
        isolate_manifest, phenotype_manifest = build_manifest(best_frame, best_sheet)
        decision = "candidate_manifest_extracted" if not phenotype_manifest.empty else "workbook_found_but_no_phenotype_rows_extracted"
        error = ""
    except Exception as exc:
        isolate_manifest = pd.DataFrame()
        phenotype_manifest = pd.DataFrame()
        decision = "workbook_found_but_manifest_extraction_failed"
        error = repr(exc)

    report = {
        "decision": decision,
        "valid_supplement_found": True,
        "supplement_path": str(supplement),
        "supplement_size_bytes": int(supplement.stat().st_size),
        "sheets": inventory.to_dict("records"),
        "selected_sheet": best_sheet,
        "candidate_isolates": int(isolate_manifest["isolate_id"].nunique()) if not isolate_manifest.empty else 0,
        "candidate_phenotype_rows": int(len(phenotype_manifest)),
        "accession_columns_detected": candidate_accession_columns(best_frame),
        "phenotype_columns_detected": candidate_phenotype_columns(best_frame),
        "error": error,
        "recommended_next_action": (
            "Audit accession columns and phenotype endpoint columns manually before downloading reads or training."
            if decision == "candidate_manifest_extracted"
            else "The workbook provides isolate/accession metadata, but no row-level MIC phenotype table was extracted. Do not train; use this as an accession manifest only unless row-level MIC labels are recovered."
        ),
    }

    Path(args.sheet_inventory).parent.mkdir(parents=True, exist_ok=True)
    Path(args.report_output).parent.mkdir(parents=True, exist_ok=True)
    inventory.to_csv(args.sheet_inventory, index=False)
    isolate_manifest.to_csv(args.isolate_output, index=False)
    phenotype_manifest.to_csv(args.phenotype_output, index=False)
    Path(args.report_output).write_text(json.dumps(report, indent=2) + "\n")
    print(f"Wrote {args.sheet_inventory}")
    print(f"Wrote {args.isolate_output}")
    print(f"Wrote {args.phenotype_output}")
    print(f"Wrote {args.report_output}")
    print(decision)


if __name__ == "__main__":
    main()
